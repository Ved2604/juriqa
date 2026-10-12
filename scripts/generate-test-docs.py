import random, sys
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import docx

random.seed(7)
OUT = sys.argv[1]
TOPICS = ["DEFINITIONS","SCOPE OF SERVICES","FEES AND PAYMENT","TERM AND TERMINATION","CONFIDENTIALITY",
 "INTELLECTUAL PROPERTY","WARRANTIES","INDEMNIFICATION","LIMITATION OF LIABILITY","INSURANCE",
 "DATA PROTECTION","FORCE MAJEURE","ASSIGNMENT","NOTICES","DISPUTE RESOLUTION","GOVERNING LAW",
 "NON-SOLICITATION","AUDIT RIGHTS","SUBCONTRACTING","GENERAL PROVISIONS"]
FILL = ("the Supplier shall perform its obligations under this Agreement with reasonable skill, care and diligence "
 "in accordance with good industry practice, all applicable laws and the reasonable instructions of the Customer as notified in writing from time to time")

def clauses(cap):
    out = []
    for n, t in enumerate(TOPICS, 1):
        subs = []
        for k in range(1, 9):
            body = f"Subject to clause {random.randint(1,20)}, {FILL}"
            if t == "LIMITATION OF LIABILITY" and k == 2:
                body = f"The Supplier's total aggregate liability arising out of or in connection with this Agreement shall not exceed {cap} in any Contract Year, whether in contract, tort (including negligence) or otherwise."
            if t == "TERM AND TERMINATION" and k == 3:
                body = "Either Party may terminate this Agreement by giving the other Party not less than ninety (90) days' written notice. Upon termination, the Customer shall pay all undisputed fees accrued up to the date of termination, including fees for work in progress, and the Supplier shall promptly return all Customer Materials in its possession."
            if k == 5:
                body = (body + ". ") * 4  # long paragraph to force splitting / page breaks
            subs.append((f"{n}.{k}", body.rstrip(". ") + "."))
        out.append((n, t, subs))
    return out

def pdf(path, cap):
    ss = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=ss["Heading2"], fontSize=12)
    p = ParagraphStyle("p", parent=ss["BodyText"], fontSize=10, leading=13, spaceAfter=6)
    def deco(c, d):
        c.saveState(); c.setFont("Helvetica", 8)
        c.drawString(72, A4[1]-40, "MASTER SERVICES AGREEMENT - CONFIDENTIAL")
        c.drawCentredString(A4[0]/2, 30, f"Page {d.page}")
        c.restoreState()
    story = [Paragraph("MASTER SERVICES AGREEMENT", ss["Title"]),
             Paragraph("This Agreement is made between Acme Trading LLC (the \"Customer\") and Beta Systems Ltd. (the \"Supplier\").", p)]
    for n, t, subs in clauses(cap):
        story.append(Paragraph(f"{n}. {t}", h))
        for num, body in subs:
            story.append(Paragraph(f"{num} {body}", p))
    SimpleDocTemplate(path, pagesize=A4).build(story, onFirstPage=deco, onLaterPages=deco)

def word(path, cap):
    d = docx.Document()
    d.add_heading("MASTER SERVICES AGREEMENT", 0)
    d.add_paragraph("This Agreement is made between Acme Trading LLC (the \"Customer\") and Beta Systems Ltd. (the \"Supplier\").")
    for n, t, subs in clauses(cap):
        d.add_heading(f"{n}. {t}", 2)
        for num, body in subs:
            para = d.add_paragraph()
            r = para.add_run(f"{num} "); r.bold = True
            para.add_run(body)
    d.save(path)

pdf(f"{OUT}/msa-v1.pdf", "AED 100,000")
word(f"{OUT}/msa-v1.docx", "AED 100,000")
pdf(f"{OUT}/msa-v2.pdf", "AED 1,000,000")
print("generated")
